import os,logging
import json
from core.mcp_instance import mcp
logger = logging.getLogger(__name__)

@mcp.tool()
async def read_file_content_tool(file_path: str) -> dict:
    """
    Smart file reader for LLM agents.
    Supports: txt, json, csv, excel, pdf, docx, pptx.
    Returns clean text or structured data.
    """

    logger.info(f"[DEBUG] read_file_content_tool called")
    logger.info(f"[DEBUG] File path: {file_path}")

    if not os.path.exists(file_path):
        logger.info("[ERROR] File not found")
        return {"error": f"File not found: {file_path}"}

    try:
        file_size = os.path.getsize(file_path)
        logger.info(f"[DEBUG] File size: {file_size} bytes")

        # Optional safety limit (10 MB)
        if file_size > 10 * 1024 * 1024:
            logger.info("[ERROR] File too large")
            return {"error": "File too large (>10MB limit)"}

        ext = os.path.splitext(file_path)[1].lower()
        logger.info(f"[DEBUG] File extension: {ext}")

        # ---------------- TEXT FILES ----------------
        if ext in [".txt", ".md", ".log"]:
            logger.info("[DEBUG] Processing TEXT file")
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()[:20000]
            return {"type": "text", "content": content}

        # ---------------- JSON ----------------
        elif ext == ".json":
            logger.info("[DEBUG] Processing JSON file")
            with open(file_path, "r", encoding="utf-8") as f:
                data = json.load(f)
            return {"type": "json", "content": data}

        # ---------------- CSV ----------------
        elif ext == ".csv":
            logger.info("[DEBUG] Processing CSV file")
            import pandas as pd
            df = pd.read_csv(file_path)

            logger.info(f"[DEBUG] CSV shape: {df.shape}")

            return {
                "type": "table",
                "columns": df.columns.tolist(),
                "rows": df.head(200).to_dict(orient="records")
            }

        # ---------------- EXCEL ----------------
        elif ext in [".xlsx", ".xls"]:
            logger.info("[DEBUG] Processing EXCEL file")
            import pandas as pd

            df_dict = pd.read_excel(file_path, sheet_name=None)
            logger.info(f"[DEBUG] Sheets found: {list(df_dict.keys())}")

            result = {}
            for sheet, df in df_dict.items():
                logger.info(f"[DEBUG] Sheet '{sheet}' shape: {df.shape}")
                result[sheet] = {
                    "columns": df.columns.tolist(),
                    "rows": df.head(200).to_dict(orient="records")
                }

            return {"type": "excel", "sheets": result}

        # ---------------- PDF ----------------
        elif ext == ".pdf":
            logger.info("[DEBUG] Processing PDF file")
            import PyPDF2

            text = ""
            with open(file_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                total_pages = len(reader.pages)
                print(f"[DEBUG] Total pages: {total_pages}")

                for i, page in enumerate(reader.pages[:20]):
                    print(f"[DEBUG] Reading page {i+1}")
                    text += page.extract_text() or ""

            return {"type": "text", "content": text[:20000]}

        # ---------------- DOCX ----------------
        elif ext == ".docx":
            logger.info("[DEBUG] Processing DOCX file")
            from docx import Document

            doc = Document(file_path)
            logger.info(f"[DEBUG] Paragraph count: {len(doc.paragraphs)}")

            text = "\n".join([p.text for p in doc.paragraphs])
            return {"type": "text", "content": text[:20000]}

        # ---------------- PPTX ----------------
        elif ext == ".pptx":
            logger.info("[DEBUG] Processing PPTX file")
            from pptx import Presentation

            prs = Presentation(file_path)
            logger.info(f"[DEBUG] Slide count: {len(prs.slides)}")

            text = []
            for i, slide in enumerate(prs.slides):
                print(f"[DEBUG] Reading slide {i+1}")
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text.append(shape.text)

            return {"type": "text", "content": "\n".join(text)[:20000]}

        # ---------------- FALLBACK ----------------
        else:
            logger.warning(f"Unsupported file type: {ext}")
            return {
                "type": "unsupported",
                "message": f"Unsupported file type: {ext}"
            }

    except Exception as e:
        logger.error(f"Exception occurred: {str(e)}")
        return {"error": str(e)}